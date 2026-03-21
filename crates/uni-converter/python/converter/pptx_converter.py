"""PPTX to Markdown converter using python-pptx."""
from pptx import Presentation
from uni_bridge import progress


def convert_pptx(path: str, options: dict) -> dict:
    progress(10, "Reading PPTX...")

    prs = Presentation(path)
    slides = prs.slides
    slide_count = len(slides)
    progress(20, f"Found {slide_count} slide(s)")

    parts = []
    has_images = False
    has_tables = False

    for i, slide in enumerate(slides):
        parts.append(f"---\n\n## Slide {i + 1}")

        if slide.shapes.title:
            title_text = slide.shapes.title.text.strip()
            if title_text:
                parts.append(f"### {title_text}")

        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        parts.append(text)

            if shape.has_table:
                has_tables = True
                table = shape.table
                rows = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows.append(cells)

                if rows:
                    parts.append("| " + " | ".join(rows[0]) + " |")
                    parts.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
                    for row in rows[1:]:
                        while len(row) < len(rows[0]):
                            row.append("")
                        parts.append("| " + " | ".join(row[: len(rows[0])]) + " |")

            if hasattr(shape, "image"):
                has_images = True

        parts.append("")

        if slide_count > 5 and (i + 1) % 5 == 0:
            pct = 20 + int(70 * (i + 1) / slide_count)
            progress(pct, f"Processing slide {i + 1}/{slide_count}...")

    markdown = "\n\n".join(parts)

    title = None
    if slides and slides[0].shapes.title:
        title = slides[0].shapes.title.text.strip() or None

    progress(95, "Finalizing...")

    return {
        "markdown": markdown,
        "title": title,
        "metadata": {
            "original_format": "pptx",
            "pages": slide_count,
            "word_count": len(markdown.split()),
            "has_images": has_images,
            "has_tables": has_tables,
            "quality": "low",
            "python_converted": True,
        },
    }
