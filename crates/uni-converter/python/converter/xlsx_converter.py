"""XLSX to Markdown converter using openpyxl."""
import openpyxl
from uni_bridge import progress


def convert_xlsx(path: str, options: dict) -> dict:
    progress(10, "Reading XLSX...")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = wb.sheetnames
    progress(20, f"Found {len(sheets)} sheet(s)")

    parts = []
    total_rows = 0

    for idx, sheet_name in enumerate(sheets):
        ws = wb[sheet_name]

        parts.append(f"## {sheet_name}\n")

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            parts.append("*(empty sheet)*\n")
            continue

        total_rows += len(rows)

        headers = [str(cell) if cell is not None else "" for cell in rows[0]]
        parts.append("| " + " | ".join(headers) + " |")
        parts.append("| " + " | ".join(["---"] * len(headers)) + " |")

        for row in rows[1:]:
            cells = [str(cell) if cell is not None else "" for cell in row]
            while len(cells) < len(headers):
                cells.append("")
            parts.append("| " + " | ".join(cells[: len(headers)]) + " |")

        parts.append("")

        pct = 20 + int(70 * (idx + 1) / len(sheets))
        progress(pct, f"Processed sheet '{sheet_name}'")

    wb.close()

    markdown = "\n".join(parts)

    progress(95, "Finalizing...")

    return {
        "markdown": markdown,
        "title": sheets[0] if sheets else None,
        "metadata": {
            "original_format": "xlsx",
            "pages": len(sheets),
            "word_count": len(markdown.split()),
            "has_images": False,
            "has_tables": True,
            "quality": "high",
            "python_converted": True,
        },
    }
